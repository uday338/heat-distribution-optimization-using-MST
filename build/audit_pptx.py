"""
audit_pptx.py - Catch text that overflows its box, before it reaches a projector.

python-pptx will happily place more text than a box can hold; PowerPoint then
spills it over whatever sits below. Nothing in the file records that, so the
only way to catch it without opening the deck is to re-measure the text with
the same fonts PowerPoint will use.

We load the real Segoe UI / Consolas TrueType files from C:\\Windows\\Fonts,
wrap each paragraph to the shape width, and compare the resulting height with
the shape height. Reports every offender with the overflow in inches.
"""
from __future__ import annotations

import glob
import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0
DPI = 96.0

FONT_FILES = {
    ("Segoe UI", False): "segoeui.ttf",
    ("Segoe UI", True): "segoeuib.ttf",
    ("Consolas", False): "consola.ttf",
    ("Consolas", True): "consolab.ttf",
}
_cache: dict = {}


def _font(name: str, bold: bool, pt: float):
    key = (name, bold, round(pt, 1))
    if key in _cache:
        return _cache[key]
    fname = FONT_FILES.get((name, bold)) or FONT_FILES[("Segoe UI", bold)]
    path = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", fname)
    px = max(1, int(round(pt * DPI / 72.0)))
    try:
        f = ImageFont.truetype(path, px)
    except OSError:
        f = ImageFont.load_default()
    _cache[key] = f
    return f


def _wrap_lines(text: str, font, max_px: float) -> int:
    """Number of lines `text` occupies when wrapped to max_px, PowerPoint-style."""
    if not text.strip():
        return 1
    lines, cur = 0, ""
    for word in text.split():
        trial = word if not cur else cur + " " + word
        if font.getlength(trial) <= max_px or not cur:
            cur = trial
        else:
            lines += 1
            cur = word
    return lines + (1 if cur else 0)


def measure(shape) -> float:
    """Estimated rendered height of a shape's text, in inches."""
    tf = shape.text_frame
    width_in = shape.width / EMU_IN
    ml = (tf.margin_left or 0) / EMU_IN
    mr = (tf.margin_right or 0) / EMU_IN
    max_px = max(8.0, (width_in - ml - mr) * DPI)

    total = (tf.margin_top or 0) / EMU_IN + (tf.margin_bottom or 0) / EMU_IN
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            total += 6 / 72.0
            continue
        pt = max((r.font.size.pt if r.font.size else 12.0) for r in runs)
        name = runs[0].font.name or "Segoe UI"
        bold = bool(runs[0].font.bold)
        text = "".join(r.text for r in runs)
        n = _wrap_lines(text, _font(name, bold, pt), max_px)
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        total += n * (pt * 1.2 * ls) / 72.0
        total += (p.space_after.pt if p.space_after else 0) / 72.0
    return total


def audit(path: str, tol: float = 0.03) -> int:
    prs = Presentation(path)
    bad = 0
    for i, slide in enumerate(prs.slides, 1):
        issues = []
        for sh in slide.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            need = measure(sh)
            have = sh.height / EMU_IN
            if need > have + tol:
                snippet = " ".join(sh.text_frame.text.split())[:52]
                # console may be cp1252; keep the report printable
                snippet = snippet.encode("ascii", "replace").decode("ascii")
                issues.append((need - have, need, have, snippet))
        if issues:
            bad += len(issues)
            print(f"  slide {i}:")
            for over, need, have, snip in sorted(issues, reverse=True):
                print(f"    +{over:.2f}in  (needs {need:.2f}, box {have:.2f})"
                      f"   {snip!r}")
    return bad


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "build/CLL798_prelim_5min.pptx"
    print(f"auditing {target}")
    n = audit(target)
    print(f"\n{n} overflowing text box(es)" if n else "\nno text overflow")
    sys.exit(1 if n else 0)
